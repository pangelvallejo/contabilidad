"""Exportación a Excel y PDF (estado de cuenta)."""
import io
from datetime import date
from decimal import Decimal

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .formato import fecha, pesos

FORMATO_PESOS = '#,##0.00;[Red]-#,##0.00'


def excel(hojas: dict) -> bytes:
    """hojas: {nombre: (encabezados, filas)}."""
    wb = Workbook()
    wb.remove(wb.active)
    for nombre, (encabezados, filas) in hojas.items():
        ws = wb.create_sheet(nombre[:31])
        ws.append(encabezados)
        for celda in ws[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="1F3A5F")
            celda.alignment = Alignment(wrap_text=True, vertical="center")
        # Una sola pasada por celda (formato al crearla): un libro diario de 20.000 líneas sale en segundos.
        for r, fila in enumerate(filas, start=2):
            for i, v in enumerate(fila, start=1):
                if isinstance(v, Decimal):
                    v = float(v)
                c = ws.cell(row=r, column=i, value=v)
                if isinstance(v, float):
                    c.number_format = FORMATO_PESOS
                elif isinstance(v, date):
                    c.number_format = "dd/mm/yyyy"
        for i, enc in enumerate(encabezados, start=1):
            letra = get_column_letter(i)
            ancho = max([len(str(enc))] + [len(str(f[i - 1])) for f in filas[:200] if i - 1 < len(f)])
            ws.column_dimensions[letra].width = min(max(10, ancho + 2), 60)
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _latin1(texto):
    texto = str(texto or "").replace("–", "-").replace("—", "-").replace("“", '"').replace("”", '"')
    return texto.encode("latin-1", "replace").decode("latin-1")


def estado_de_cuenta_pdf(empresa: dict, cliente, filas, total, corte: date) -> bytes:
    """filas: lista de dicts con numero, fecha, vencimiento, total, reteiva, abonos, saldo, dias."""
    pdf = FPDF(orientation="P", unit="mm", format="Letter")
    pdf.set_auto_page_break(True, margin=15)
    pdf.add_page()
    pdf.set_font("helvetica", "B", 14)
    pdf.cell(0, 7, _latin1(empresa["nombre"]), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("helvetica", "", 9)
    pdf.cell(0, 5, _latin1(f"NIT {empresa['nit']}  {empresa.get('direccion') or ''}  {empresa.get('ciudad') or ''}"),
             new_x="LMARGIN", new_y="NEXT")
    if empresa.get("email") or empresa.get("telefono"):
        pdf.cell(0, 5, _latin1(f"{empresa.get('email') or ''}  {empresa.get('telefono') or ''}"),
                 new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    pdf.set_font("helvetica", "B", 12)
    pdf.cell(0, 7, "ESTADO DE CUENTA", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("helvetica", "", 10)
    pdf.cell(0, 6, _latin1(f"Cliente: {cliente.nombre}   NIT {cliente.nit_completo}"), new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 6, _latin1(f"Fecha de corte: {fecha(corte)}"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)

    columnas = [("Factura", 24, "L"), ("Fecha", 21, "C"), ("Vence", 21, "C"), ("Valor", 28, "R"),
                ("ReteIVA", 22, "R"), ("Abonos y NC", 28, "R"), ("Saldo", 28, "R"), ("Días venc.", 18, "R")]
    pdf.set_font("helvetica", "B", 8)
    pdf.set_fill_color(31, 58, 95)
    pdf.set_text_color(255)
    for titulo, ancho, al in columnas:
        pdf.cell(ancho, 7, _latin1(titulo), border=0, align=al, fill=True)
    pdf.ln()
    pdf.set_text_color(0)
    pdf.set_font("helvetica", "", 8)
    for i, f in enumerate(filas):
        pdf.set_fill_color(242, 245, 249)
        valores = [f["numero"], fecha(f["fecha"]), fecha(f["vencimiento"]), pesos(f["total"]),
                   pesos(f["reteiva"]), pesos(f["abonos"]), pesos(f["saldo"]),
                   str(f["dias"]) if f["dias"] > 0 else "-"]
        for (titulo, ancho, al), v in zip(columnas, valores):
            pdf.cell(ancho, 6, _latin1(v), align=al, fill=i % 2 == 0)
        pdf.ln()
    pdf.set_font("helvetica", "B", 10)
    pdf.ln(2)
    pdf.cell(sum(c[1] for c in columnas[:6]), 7, "TOTAL A PAGAR", align="R")
    pdf.cell(columnas[6][1], 7, _latin1(pesos(total)), align="R")
    pdf.ln(12)
    pdf.set_font("helvetica", "", 8)
    pdf.multi_cell(0, 4, _latin1(
        "Si ya realizó el pago, por favor haga caso omiso de este estado de cuenta y envíenos el soporte. "
        "Si nos practicó retención de IVA, le agradecemos remitir el certificado correspondiente."))
    return bytes(pdf.output())


class _Informe(FPDF):
    def __init__(self, empresa, titulo, subtitulo="", orientacion="P"):
        super().__init__(orientation=orientacion, unit="mm", format="Letter")
        self.empresa, self.titulo, self.subtitulo = empresa, titulo, subtitulo
        self.set_auto_page_break(True, margin=15)

    def header(self):
        self.set_font("helvetica", "B", 12)
        self.cell(0, 6, _latin1(self.empresa.get("nombre", "")), new_x="LMARGIN", new_y="NEXT")
        self.set_font("helvetica", "", 8)
        self.cell(0, 4, _latin1(f"NIT {self.empresa.get('nit', '')}  {self.empresa.get('ciudad') or ''}"),
                  new_x="LMARGIN", new_y="NEXT")
        self.ln(2)
        self.set_font("helvetica", "B", 13)
        self.cell(0, 7, _latin1(self.titulo), new_x="LMARGIN", new_y="NEXT")
        if self.subtitulo:
            self.set_font("helvetica", "", 9)
            self.cell(0, 5, _latin1(self.subtitulo), new_x="LMARGIN", new_y="NEXT")
        self.ln(3)

    def footer(self):
        self.set_y(-12)
        self.set_font("helvetica", "I", 7)
        self.cell(0, 5, _latin1(f"Página {self.page_no()} · generado el {date.today():%d/%m/%Y}"), align="R")


def _empresa_dict(session):
    from .contab import config as cfg
    emp = {k: cfg(session, f"empresa_{k}", "") for k in ("nombre", "direccion", "ciudad", "email", "telefono")}
    emp["nit"] = f"{cfg(session, 'empresa_nit', '')}-{cfg(session, 'empresa_dv', '')}"
    return emp


def informe_pdf(empresa, titulo, subtitulo, columnas, filas, orientacion="P", notas=None) -> bytes:
    """Tabla genérica. columnas: [(título, ancho_mm, alineación)]; filas: listas de celdas; una fila puede ser
    un dict {"celdas": [...], "estilo": "grupo"|"total"|"res"} para resaltarla."""
    pdf = _Informe(empresa, titulo, subtitulo, orientacion)
    pdf.add_page()
    ancho_total = sum(c[1] for c in columnas)

    def encabezado():
        pdf.set_font("helvetica", "B", 8)
        pdf.set_fill_color(31, 58, 95)
        pdf.set_text_color(255)
        for t, w, al in columnas:
            pdf.cell(w, 6, _latin1(t), align=al, fill=True)
        pdf.ln()
        pdf.set_text_color(0)

    encabezado()
    for i, fila in enumerate(filas):
        estilo = ""
        celdas = fila
        if isinstance(fila, dict):
            celdas, estilo = fila["celdas"], fila.get("estilo", "")
        if pdf.get_y() > pdf.h - 25:
            pdf.add_page()
            encabezado()
        pdf.set_font("helvetica", "B" if estilo else "", 8)
        if estilo == "total":
            pdf.set_fill_color(220, 226, 235)
        else:
            pdf.set_fill_color(242, 245, 249)
        relleno = estilo in ("total", "grupo") or (not estilo and i % 2 == 0)
        for (t, w, al), v in zip(columnas, celdas):
            texto = pesos(v) if isinstance(v, Decimal) else (fecha(v) if isinstance(v, date) else str(v))
            pdf.cell(w, 5.5, _latin1(texto)[:int(w / 1.7)], align=al, fill=relleno)
        pdf.ln()
    if notas:
        pdf.ln(4)
        pdf.set_font("helvetica", "", 8)
        for n in notas:
            pdf.multi_cell(ancho_total, 4, _latin1(n))
    return bytes(pdf.output())


def recibo_de_caja_pdf(empresa, rec) -> bytes:
    """Comprobante para el cliente de un pago recibido y las facturas que cubrió."""
    pdf = _Informe(empresa, f"Recibo de caja No. {rec.id}", f"Fecha: {fecha(rec.fecha)}")
    pdf.add_page()
    pdf.set_font("helvetica", "", 10)
    pdf.cell(0, 6, _latin1(f"Recibido de: {rec.cliente.nombre}   NIT {rec.cliente.nit_completo}"), new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 6, _latin1(f"Valor recibido: {pesos(rec.valor)}   Cuenta: {rec.banco.nombre}"
                           f"{'   Ref. ' + rec.referencia if rec.referencia else ''}"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    columnas = [("Factura", 40, "L"), ("Fecha", 30, "C"), ("Valor factura", 40, "R"), ("Abono", 40, "R")]
    pdf.set_font("helvetica", "B", 9)
    pdf.set_fill_color(31, 58, 95)
    pdf.set_text_color(255)
    for t, w, al in columnas:
        pdf.cell(w, 6, t, align=al, fill=True)
    pdf.ln()
    pdf.set_text_color(0)
    pdf.set_font("helvetica", "", 9)
    for a in rec.aplicaciones:
        for (t, w, al), v in zip(columnas, [a.documento.numero, fecha(a.documento.fecha), pesos(a.documento.total),
                                            pesos(a.valor)]):
            pdf.cell(w, 6, _latin1(v), align=al)
        pdf.ln()
    if rec.sin_aplicar > 0:
        pdf.cell(110, 6, "Anticipo (sin aplicar a facturas)", align="R")
        pdf.cell(40, 6, _latin1(pesos(rec.sin_aplicar)), align="R")
        pdf.ln()
    pdf.ln(10)
    pdf.set_font("helvetica", "", 8)
    pdf.multi_cell(0, 4, _latin1("Este recibo acredita el pago recibido. La factura electrónica correspondiente fue "
                                 "emitida a través de la DIAN."))
    return bytes(pdf.output())


def cotizacion_pdf(empresa, cot) -> bytes:
    """Propuesta de honorarios para el cliente, con las líneas, IVA, total y condiciones."""
    pdf = _Informe(empresa, f"Propuesta de honorarios {cot.numero}",
                   f"Fecha: {fecha(cot.fecha)} · válida hasta el {fecha(cot.vence)}")
    pdf.add_page()
    pdf.set_font("helvetica", "", 10)
    pdf.cell(0, 6, _latin1(f"Para: {cot.cliente.nombre}   NIT {cot.cliente.nit_completo}"), new_x="LMARGIN", new_y="NEXT")
    if cot.cliente.email or cot.cliente.direccion:
        pdf.cell(0, 6, _latin1(f"{cot.cliente.direccion or ''}  {cot.cliente.email or ''}"), new_x="LMARGIN", new_y="NEXT")
    if cot.asunto:
        pdf.cell(0, 6, _latin1(f"Asunto: {cot.asunto.nombre}"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf.set_font("helvetica", "B", 11)
    pdf.multi_cell(0, 6, _latin1(cot.titulo))
    pdf.ln(2)
    columnas = [("Concepto", 115, "L"), ("Honorarios", 30, "R"), ("IVA", 25, "R"), ("Total", 25, "R")]
    pdf.set_font("helvetica", "B", 9)
    pdf.set_fill_color(31, 58, 95)
    pdf.set_text_color(255)
    for t, w, al in columnas:
        pdf.cell(w, 6, t, align=al, fill=True)
    pdf.ln()
    pdf.set_text_color(0)
    pdf.set_font("helvetica", "", 9)
    for l in cot.lineas:
        y = pdf.get_y()
        pdf.multi_cell(115, 5, _latin1(l.descripcion), new_x="RIGHT", new_y="TOP")
        alto = max(pdf.get_y() - y, 5)
        pdf.set_xy(pdf.l_margin + 115, y)
        pdf.cell(30, alto, _latin1(pesos(l.valor)), align="R")
        pdf.cell(25, alto, _latin1(pesos(l.iva)), align="R")
        pdf.cell(25, alto, _latin1(pesos(l.valor + l.iva)), align="R")
        pdf.ln(alto)
    pdf.ln(2)
    pdf.set_font("helvetica", "B", 10)
    for etiqueta, v in (("Subtotal honorarios", cot.subtotal), ("IVA 19 %", cot.iva), ("TOTAL", cot.total)):
        pdf.cell(145, 6, etiqueta, align="R")
        pdf.cell(50, 6, _latin1(pesos(v)), align="R")
        pdf.ln()
    if cot.condiciones:
        pdf.ln(4)
        pdf.set_font("helvetica", "B", 9)
        pdf.cell(0, 5, "Condiciones", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("helvetica", "", 9)
        pdf.multi_cell(0, 4.5, _latin1(cot.condiciones))
    pdf.ln(6)
    pdf.set_font("helvetica", "", 8)
    pdf.multi_cell(0, 4, _latin1("Los honorarios se facturarán electrónicamente a través de la DIAN una vez aceptada "
                                 "esta propuesta. Régimen SIMPLE de Tributación: no practicar retención en la fuente "
                                 "a título de renta (art. 911 E.T.)."))
    return bytes(pdf.output())
